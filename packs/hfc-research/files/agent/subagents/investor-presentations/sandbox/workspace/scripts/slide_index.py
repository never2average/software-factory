#!/usr/bin/env python3
"""Index a deck slide by slide: title, unit, periods, section guess.

  python3 /workspace/scripts/slide_index.py /workspace/in/deck.pdf  > /workspace/out/slide-index.json
  python3 /workspace/scripts/slide_index.py /workspace/in/deck.pptx > /workspace/out/slide-index.json
  python3 /workspace/scripts/slide_index.py --text-file slides.txt   (slides separated by a form feed, \\f)

For every slide: the title (PPTX title placeholder, else the largest text near the top, else the first real line),
the unit the slide names (null when none OR more than one: see units_found), the period tokens on it, whether it
mixes period kinds, a section guess from references/section-keywords.json (null on a tie or a weak score), whether it
sits after the Appendix divider, and whether it is a restructured-book slide (excluded: nothing is extracted from it).

The output validates against schemas/slide-index.schema.json. pdfplumber / python-pptx are imported lazily.
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import argparse, re
from collections import Counter
from finlib import units, schema, pdfdoc
from iplib import emit, fail, load_table, load_schema, find_phrases, find_periods, find_units

_USD = re.compile(r"US\s*\$|USD|\$\s*\d|\bdollars?\b", re.I)
_NOISE = [
    re.compile(r"^\s*(?:page\s*)?\d{1,3}\s*(?:of\s*\d{1,3})?\s*$", re.I),                     # page number
    re.compile(r"^\s*\d{1,3}\s*[|•·-]\s*\S.*$"),                                                # '12 | Investor Presentation'
    re.compile(r"^.*[|•·]\s*\d{1,3}\s*$"),                                                      # 'Investor Presentation | 12'
    re.compile(r"^\s*[\(\[]?\s*(?:all\s+)?(?:figures|amounts?|values)?\s*(?:are\s+)?(?:in\s+)?(?:₹|rs\.?|inr)\s*(?:in\s+)?(?:crores?|cr\.?|lakhs?|millions?|mn|billions?|bn)\s*(?:unless.*)?[\)\]]?\s*$", re.I),
    re.compile(r"^\s*(?:private\s+(?:and|&)\s+confidential|strictly\s+confidential|confidential)\s*$", re.I),
    re.compile(r"^\s*(?:www\.|https?://)\S+\s*$", re.I),
]


def is_noise(line):
    s = line.strip()
    if len(re.findall(r"[A-Za-z]", s)) < 3:
        return True
    return any(rx.match(s) for rx in _NOISE)


def title_from_lines(lines):
    """First line that is not a page number, footer, unit banner or URL. -> (title|None, method)"""
    for ln in lines:
        if not is_noise(ln) and len(ln.strip()) <= 140:
            return " ".join(ln.split()), "first_line"
    return None, "none"


def title_from_sized_lines(lines):
    """lines: [(text, font_size, top_fraction 0..1)]. The largest text in the top 40% of the slide; a slide with at most
    three real lines (a section divider) takes the largest anywhere. Ties go to the topmost. -> (title|None, method)"""
    real = [(t, s, y) for t, s, y in lines if not is_noise(t)]
    if not real:
        return None, "none"
    pool = real if len(real) <= 3 else [r for r in real if r[2] <= 0.40]
    if not pool:
        t, m = title_from_lines([t for t, _, _ in sorted(real, key=lambda r: r[2])])
        return t, m
    big = max(s for _, s, _ in pool)
    top = sorted([r for r in pool if r[1] >= big - 0.5], key=lambda r: r[2])
    # a title wrapped over two lines: join the next line when it has the same size and sits directly below
    title = top[0][0]
    if len(top) > 1 and top[1][2] - top[0][2] < 0.09:
        title = title + " " + top[1][0]
    return " ".join(title.split()), "largest_text"


def score_sections(title, text, table):
    tw, bw = table["title_weight"], table["body_weight"]
    out = []
    for name, spec in table["sections"].items():
        in_title = find_phrases(title or "", spec["phrases"])
        in_body = [p for p in find_phrases(text or "", spec["phrases"]) if p not in in_title]
        score = tw * len(in_title) + bw * len(in_body)
        if score:
            out.append({"section": name, "score": score, "matched": in_title + in_body})
    out.sort(key=lambda c: (-c["score"], c["section"]))
    return out


def pick_section(cands, table):
    if not cands or cands[0]["score"] < table["min_score"]:
        return None, 0
    if len(cands) > 1 and cands[1]["score"] == cands[0]["score"]:
        return None, cands[0]["score"]            # a tie is reported, never broken by guessing
    return cands[0]["section"], cands[0]["score"]


def is_appendix_divider(title, line_count, table):
    t = re.sub(r"[^a-z ]", " ", (title or "").lower()).split()
    t = " ".join(t)
    if not t:
        return False
    if t in table["appendix_divider_titles"]:
        return True
    return line_count <= 4 and any(t.startswith(d + " ") for d in ("appendix", "annexure", "annexures", "appendices"))


def index_slide(n, text, title, method, table, keep_text=True):
    ufound = find_units(text)
    per = [{"label": p["label"], "period": p["period"], "kind": p["kind"]} for p in find_periods(text)]
    seen, uniq = set(), []
    for p in per:
        k = (p["period"], p["kind"], p["label"].lower())
        if k not in seen:
            seen.add(k); uniq.append(p)
    flow_kinds = {p["kind"] for p in uniq if p["kind"] in ("quarter", "cumulative", "year", "trailing")}
    cands = score_sections(title, text, table)
    section, score = pick_section(cands, table)
    if n == 1 and section is None and any(c["section"] == "cover" for c in cands):
        section, score = "cover", next(c["score"] for c in cands if c["section"] == "cover")   # slide 1 naming itself a presentation
    restructured_title = bool(find_phrases(title or "", table["sections"]["restructured_book"]["phrases"]))
    restructured_body = bool(find_phrases(text or "", table["sections"]["restructured_book"]["phrases"]))
    if restructured_title:
        section = "restructured_book"
    chars = len((text or "").strip())
    row = {"slide": n, "title": title, "title_method": method, "unit": units.detect_unit(text) if len(ufound) <= 1 else None,
           "units_found": ufound, "usd_present": bool(_USD.search(text or "")), "periods": uniq, "mixed_periods": len(flow_kinds) > 1,
           "section": section, "section_score": score, "section_candidates": cands[:4], "in_appendix": False,
           "excluded": restructured_title, "restructured_mentions": restructured_body, "char_count": chars, "image_only": chars < 40}
    if keep_text:
        row["text"] = text or ""
    return row


def finish(document, kind, rows, table):
    start = None
    for r in rows:
        lines = [l for l in (r.get("text") or r.get("_text") or "").splitlines() if l.strip()]
        if start is None and is_appendix_divider(r["title"], len(lines), table):
            start = r["slide"]
        r["in_appendix"] = start is not None and r["slide"] >= start
        r.pop("_text", None)
    warnings = []
    img = [r["slide"] for r in rows if r["image_only"]]
    if img: warnings.append(f"slides with little or no text (section dividers, charts or images): {img}")
    multi = [r["slide"] for r in rows if len(r["units_found"]) > 1]
    if multi: warnings.append(f"slides naming more than one unit, unit left null: {multi}")
    ex = [r["slide"] for r in rows if r["excluded"]]
    if ex: warnings.append(f"restructured-book slides, excluded by the analysts' rulebook: {ex}")
    if start is None: warnings.append("no Appendix/Annexure divider slide found; in_appendix is false everywhere")
    c = Counter(r["unit"] for r in rows if r["unit"])
    top = c.most_common(2)
    deck_unit = top[0][0] if top and (len(top) == 1 or top[0][1] > top[1][1]) else None
    out = {"document": document, "kind": kind, "slide_count": len(rows), "appendix_starts_at": start, "deck_unit": deck_unit,
           "warnings": warnings, "slides": rows}
    problems = schema.validate(out, load_schema("slide-index.schema.json"))
    if problems:
        fail("slide index does not match schemas/slide-index.schema.json:\n  " + "\n  ".join(problems[:20]), 5)
    return out


def _pdf_sized_lines(page):
    words = page.extract_words(extra_attrs=["size"]) or []
    rows = {}
    for w in words:
        rows.setdefault(round(w["top"] / 3.0), []).append(w)
    h = float(page.height or 1)
    out = []
    for k in sorted(rows):
        ws = sorted(rows[k], key=lambda w: w["x0"])
        out.append((" ".join(w["text"] for w in ws), max(float(w.get("size") or 0) for w in ws), min(w["top"] for w in ws) / h))
    return out


def index_pdf(path, table, keep_text):
    import pdfplumber  # lazy
    rows = []
    with pdfplumber.open(path) as pdf:
        for n, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            try:
                title, method = title_from_sized_lines(_pdf_sized_lines(page))
            except Exception:
                title, method = title_from_lines(text.splitlines())
            r = index_slide(n, text, title, method, table, keep_text)
            if not keep_text: r["_text"] = text
            rows.append(r)
    return rows


def index_pptx(path, table, keep_text):
    from pptx import Presentation  # lazy (python-pptx)
    from pptx.util import Pt
    rows = []
    prs = Presentation(path)
    height = float(prs.slide_height or 1)
    for n, slide in enumerate(prs.slides, 1):
        parts, sized, title, method = [], [], None, "none"
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame and slide.shapes.title.text_frame.text.strip():
            title, method = " ".join(slide.shapes.title.text_frame.text.split()), "placeholder"

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6 and hasattr(sh, "shapes"):   # group
                    walk(sh.shapes); continue
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                    for para in sh.text_frame.paragraphs:
                        t = "".join(r.text for r in para.runs).strip()
                        if not t: continue
                        parts.append(t)
                        sizes = [r.font.size.pt for r in para.runs if r.font.size is not None]
                        sized.append((t, max(sizes) if sizes else 0.0, float(sh.top or 0) / height))
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        parts.append("  ".join(c.text.strip() for c in row.cells))
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    parts.append("[chart]")
        walk(slide.shapes)
        text = "\n".join(parts)
        if title is None:
            title, method = title_from_sized_lines(sized) if any(s for _, s, _ in sized) else title_from_lines(parts)
        r = index_slide(n, text, title, method, table, keep_text)
        if not keep_text: r["_text"] = text
        rows.append(r)
    return rows


def index_text(text, table, keep_text):
    rows = []
    for n, chunk in enumerate(text.split("\f"), 1):
        title, method = title_from_lines(chunk.splitlines())
        r = index_slide(n, chunk, title, method, table, keep_text)
        if not keep_text: r["_text"] = chunk
        rows.append(r)
    return rows


_SYNTH = "\f".join([
    "Example Housing Finance Ltd\nInvestor Presentation\nQ2 FY26",
    "Disclaimer\nThis presentation contains forward-looking statements.\n2",
    "3 | Investor Presentation Q2 FY26\nKey Highlights – Q2 FY26\n(₹ in crore)\nAUM 12,345 up 18% YoY\nDisbursements 1,050\nBranches 215 Employees 3,410\nGNPA 1.42%",
    "AUM and Disbursement trend\n₹ crore\nQ2 FY25 Q1 FY26 Q2 FY26 H1 FY26\nDisbursements 890 960 1,050 2,010\nAUM 10,460 11,870 12,345",
    "Borrowing profile and ALM\nINR bn (US$ mn in brackets)\nTerm loans 62% NCDs 18% NHB refinance 12%\nLiquidity position as on 30 September 2025",
    "Yield, cost of funds and spread\nYield 11.6% Cost of funds 8.1% Spread 3.5% NIM 4.2%",
    "Asset quality\nGross Stage 3 1.42% Net Stage 3 0.98% Stage 3 PCR 31%\nCollection efficiency 99.1%",
    "Appendix",
    "Details of loans transferred through direct assignment\n₹ crore\nQ2 FY26 H1 FY26\nLoans assigned during the quarter 180 310\nPortfolio buyout Nil Nil",
    "Restructured book\n₹ crore\nOTR 2.0 outstanding 95",
    "Glossary\nAUM: on-book loans plus assigned and co-lent loans\nSpread: yield on average loans minus cost of average borrowings",
    "7",
    "Quarter in brief\nYield and spread held; GNPA and collection efficiency steady",
])


def _self_test():
    table = load_table("section-keywords.json")
    checks = []
    def ok(name, cond, detail=None):
        checks.append({"check": name, "ok": bool(cond), **({"detail": detail} if not cond else {})})

    out = finish("synthetic", "text", index_text(_SYNTH, table, True), table)   # also proves the output matches the schema
    s = {r["slide"]: r for r in out["slides"]}
    ok("13 slides indexed", out["slide_count"] == 13)
    ok("cover and disclaimer", s[1]["section"] == "cover" and s[2]["section"] == "disclaimer", (s[1]["section"], s[2]["section"]))
    ok("title skips the '3 | Investor Presentation' footer line", s[3]["title"] == "Key Highlights – Q2 FY26", s[3]["title"])
    ok("highlights slide: section, unit crore, period Q2 FY26", s[3]["section"] == "highlights" and s[3]["unit"] == "crore" and
       [p["period"] for p in s[3]["periods"]][:1] == ["Q2 FY26"], s[3])
    ok("trend slide: aum_disbursements, mixed periods (quarter + H1)", s[4]["section"] == "aum_disbursements" and s[4]["mixed_periods"] is True, s[4])
    ok("borrowings slide: two units -> unit null, USD seen", s[5]["unit"] is None and s[5]["units_found"] == ["million", "billion"] and s[5]["usd_present"], s[5])
    ok("as-at date parsed to its quarter", any(p["kind"] == "as_at" and p["period"] == "Q2 FY26" for p in s[5]["periods"]), s[5]["periods"])
    ok("margins and asset quality sections", s[6]["section"] == "margins_spreads" and s[7]["section"] == "asset_quality", (s[6]["section"], s[7]["section"]))
    ok("no unit on a percent-only slide", s[6]["unit"] is None and s[6]["units_found"] == [])
    ok("appendix divider found at 8; later slides flagged, earlier not", out["appendix_starts_at"] == 8 and s[9]["in_appendix"] and not s[7]["in_appendix"])
    ok("sell-down slide in appendix", s[9]["section"] == "appendix" and s[9]["in_appendix"], s[9]["section_candidates"])
    ok("restructured-book slide is excluded", s[10]["excluded"] and s[10]["section"] == "restructured_book", s[10])
    ok("glossary", s[11]["section"] == "glossary", s[11]["section_candidates"])
    ok("page-number-only slide: no title, image_only", s[12]["title"] is None and s[12]["title_method"] == "none" and s[12]["image_only"])
    ok("tie between sections -> section null with candidates", s[13]["section"] is None and len(s[13]["section_candidates"]) >= 2, s[13]["section_candidates"])
    ok("warnings name the image-only, multi-unit and excluded slides", len(out["warnings"]) == 3, out["warnings"])
    ok("deck unit is the most common slide unit", out["deck_unit"] == "crore", out["deck_unit"])

    t, m = title_from_sized_lines([("Example Housing Finance Ltd", 9, 0.02), ("Product and customer mix", 28, 0.08), ("Individual housing 71%", 14, 0.3),
                                   ("LAP 19%", 14, 0.4), ("71%", 40, 0.55), ("12", 9, 0.97)])
    ok("largest text in the top 40% wins over a bigger KPI number lower down", (t, m) == ("Product and customer mix", "largest_text"), (t, m))
    t, m = title_from_sized_lines([("Appendix", 44, 0.5), ("14", 9, 0.97)])
    ok("divider slide: centred title still found", t == "Appendix", t)
    t, m = title_from_sized_lines([("Diversified borrowing profile with", 26, 0.06), ("a granular lender base", 26, 0.12), ("Term loans 62%", 12, 0.3),
                                   ("NCD 18%", 12, 0.35), ("NHB 12%", 12, 0.4)])
    ok("two-line title is joined", t == "Diversified borrowing profile with a granular lender base", t)
    ok("no real lines -> none", title_from_sized_lines([("12", 9, 0.9)]) == (None, "none"))
    ok("unit banner is not a title", title_from_lines(["(₹ in crore)", "Rs. in Lakhs", "Capital adequacy"]) == ("Capital adequacy", "first_line"))

    out2 = finish("synthetic", "text", index_text("Key highlights\n₹ crore\nAUM 100", table, False), table)
    ok("--no-text drops the text and reports the missing appendix divider", "text" not in out2["slides"][0] and any("no Appendix" in w for w in out2["warnings"]))
    return checks


def main():
    ap = argparse.ArgumentParser(description="Per-slide index of a deck: title, unit, periods, section guess. JSON on stdout.")
    ap.add_argument("file", nargs="?", help="a .pdf or .pptx in the sandbox")
    ap.add_argument("--text-file", help="plain text with slides separated by a form feed (\\f); '-' for stdin")
    ap.add_argument("--no-text", action="store_true", help="leave the slide text out of the index (smaller output; find_metric_slides.py then matches titles only)")
    ap.add_argument("--self-test", action="store_true", help="run built-in cases (no input files, no pdfplumber) and exit 0 or 1")
    a = ap.parse_args()
    if a.self_test:
        checks = _self_test()
        bad = [c for c in checks if not c["ok"]]
        emit({"script": "slide_index.py", "passed": len(checks) - len(bad), "failed": len(bad), "failures": bad})
        sys.exit(1 if bad else 0)
    table = load_table("section-keywords.json")
    keep = not a.no_text
    if a.text_file:
        text = sys.stdin.read() if a.text_file == "-" else open(a.text_file, encoding="utf-8").read()
        emit(finish(a.text_file, "text", index_text(text, table, keep), table)); return
    if not a.file:
        fail("give a .pdf or .pptx path, --text-file, or --self-test")
    if not os.path.isfile(a.file):
        fail(f"no such file: {a.file}")
    kind = pdfdoc.sniff(a.file)
    try:
        if kind == "pdf":
            rows = index_pdf(a.file, table, keep); k = "pdf"
        elif kind == "zip-office":
            rows = index_pptx(a.file, table, keep); k = "pptx"
        else:
            fail(f"{a.file} is neither a PDF nor a PPTX (looks like: {kind}); run detect_content_type.py and report it")
    except ImportError as x:
        fail(f"a document library is missing in this sandbox ({x}); cannot index the deck", 3)
    except SystemExit:
        raise
    except Exception as x:
        fail(f"could not read {a.file} ({type(x).__name__}: {x}); report the file as unreadable", 4)
    emit(finish(os.path.basename(a.file), k, rows, table))


if __name__ == "__main__":
    main()
