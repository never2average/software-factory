#!/usr/bin/env python3
"""kpis.jsonl -> the analysts' KPI workbook.

  python3 /workspace/scripts/build_kpi_workbook.py /workspace/out/kpis-all.jsonl --xlsx /workspace/out/example-hfl-kpis.xlsx
  python3 /workspace/scripts/build_kpi_workbook.py /workspace/out/kpis-all.jsonl --customer example-hfl --name "Example Housing Finance Ltd" > spec.json
  python3 /workspace/scripts/build_kpi_workbook.py --self-test

Layout
  one sheet per company   rows = the 27 KPIs grouped by category in rulebook order; columns = Category, KPI, Unit, one column per
                          quarter (oldest first), Notes. A not_found cell reads "not found"; a KPI never extracted for a quarter is blank.
                          The Notes column lists the footnote references of the row, e.g. "Q2 FY26 [3] needs review".
  "Footnotes"             one row per footnote: Ref, Company, KPI, Period, Status, Value period, Footnote.
  "Citations"             every cell's source, document and page or slide (each value must be traceable).
The input is validated first with validate_kpis; an invalid file builds nothing (exit 1). When a cell appears more than once the row
with the latest extracted_at is used. Input rows are the whole history to show (the data room's kpis.jsonl plus the new batch).
The spec JSON ({"workbook": {"filename", "sheets": [{"name", "columns", "rows", "cell_status"}]}}) is always printed; the .xlsx is written
when --xlsx is given and openpyxl can be imported (`xlsx.written` says which). Afterwards: python3 /root/fmt_xlsx.py <file.xlsx>.
"""
import os, sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import argparse, json, re, tempfile
from finlib import periods, schema
import kpi_catalog as cat
import validate_kpis

NOT_FOUND_TEXT = "not found"
STATUS_WORDS = {"needs_review": "needs review", "carried_forward": "carried forward", "not_found": "not found", "ok": ""}
FILL = {"needs_review": "FFE8A3", "carried_forward": "DCE6F5", "not_found": "EEEEEE"}


def sheet_name(text, used):
    s = re.sub(r"[\[\]\:\*\?\/\\]", " ", text).strip().strip("'")[:31] or "Company"
    base, i = s, 2
    while s.lower() in used:
        suffix = f" ({i})"; s = base[:31 - len(suffix)] + suffix; i += 1
    used.add(s.lower()); return s


def _pkey(p):
    n = periods.normalise(p); return n["fy"] * 10 + n["quarter"]


def build(rows, customer=None, names=None):
    """rows: [obj] (already valid) -> workbook spec"""
    names = names or {}
    latest = {}
    for o in rows:
        if customer and o["customer_id"] != customer: continue
        key = (o["customer_id"], o["period"], o["kpi"])
        if key not in latest or o["extracted_at"] >= latest[key]["extracted_at"]: latest[key] = o
    companies = sorted({k[0] for k in latest})
    used = {"footnotes", "citations"}
    sheets, foot_rows, cite_rows = [], [], []
    ref = 0
    for cid in companies:
        quarters = sorted({k[1] for k in latest if k[0] == cid}, key=_pkey)
        columns = ["Category", "KPI", "Unit"] + quarters + ["Notes"]
        body, cell_status = [], []
        for e in cat.CATALOG:
            line, notes = [e["category"], e["label"], e["unit"]], []
            for q in quarters:
                o = latest.get((cid, q, e["key"]))
                if o is None: line.append(None); continue
                line.append(NOT_FOUND_TEXT if o["status"] == "not_found" else o["value"])
                if o["status"] != "ok": cell_status.append({"row": len(body) + 2, "column": columns.index(q) + 1, "status": o["status"]})
                if o.get("footnote") or o["status"] != "ok":
                    ref += 1
                    notes.append(f"{q} [{ref}]" + (f" {STATUS_WORDS[o['status']]}" if o["status"] != "ok" else ""))
                    foot_rows.append([ref, names.get(cid, cid), e["label"], q, o["status"], o.get("value_period"), o.get("footnote", "")])
                cite_rows.append([names.get(cid, cid), e["label"], q, o["value"], e["unit"], o["basis"], o["source"], o.get("document"),
                                  None if o.get("page_or_slide") is None else str(o["page_or_slide"]), o.get("definition"), o["extracted_at"]])
            line.append("; ".join(notes) or None)
            body.append(line)
        sheets.append({"name": sheet_name(names.get(cid, cid), used), "customer_id": cid, "columns": columns, "rows": body, "cell_status": cell_status})
    sheets.append({"name": "Footnotes", "columns": ["Ref", "Company", "KPI", "Period", "Status", "Value period", "Footnote"], "rows": foot_rows, "cell_status": []})
    sheets.append({"name": "Citations", "columns": ["Company", "KPI", "Period", "Value", "Unit", "Basis", "Source", "Document", "Page or slide", "Definition", "Extracted at"],
                   "rows": cite_rows, "cell_status": []})
    fname = (companies[0] if len(companies) == 1 else "hfc") + "-kpis.xlsx"
    return {"filename": fname, "companies": companies, "sheets": sheets}


def write_xlsx(spec, path):
    """-> (written, reason). openpyxl is imported here, lazily."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
    except ImportError:
        return False, "openpyxl is not importable in this environment; the spec was still produced. In the sandbox it is installed by sandbox.ts: report this as a sandbox fault."
    wb = Workbook(); wb.remove(wb.active)
    for sh in spec["sheets"]:
        ws = wb.create_sheet(sh["name"])
        ws.append(sh["columns"])
        for r in sh["rows"]: ws.append(r)
        for c in sh["cell_status"]:
            if c["status"] in FILL: ws.cell(row=c["row"], column=c["column"]).fill = PatternFill("solid", fgColor=FILL[c["status"]])
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return True, None


def run(path, customer=None, names=None, xlsx=None):
    rows, problems = schema.read_jsonl(path)
    rep = validate_kpis.validate_rows(rows, problems)
    if not rep["valid"]:
        return {"ok": False, "problem": "the input fails validate_kpis; nothing was built", "errors": rep["errors"][:50]}
    objs = [o for _, o in rows]
    if customer and not any(o["customer_id"] == customer for o in objs):
        return {"ok": False, "problem": f"no rows for customer {customer!r}", "errors": []}
    spec = build(objs, customer, names)
    out = {"ok": True, "workbook": spec, "flags": rep["flags"], "xlsx": {"requested": bool(xlsx), "written": False, "path": xlsx, "reason": None if xlsx else "no --xlsx given"}}
    if xlsx:
        w, why = write_xlsx(spec, xlsx); out["xlsx"].update(written=w, reason=why)
    return out


def _self_test():
    fails = []
    def eq(name, got, want):
        if got != want: fails.append(f"{name}: got {got!r}, want {want!r}")
    def row(cid, period, kpi, value, ts="2026-09-18T10:00:00Z", **kw):
        e = cat.BY_KEY[kpi]
        r = {"customer_id": cid, "extracted_at": ts, "kpi": kpi, "category": e["category"], "value": value, "unit": e["unit"], "period": period,
             "basis": "standalone", "source": "computed" if e["source_pref"] == "computed" else e["source_pref"],
             "document": f"Customers/{cid}/filings/lodr/results.pdf", "page_or_slide": 3, "status": "ok", "footnote": ""}
        r.update(kw); return r
    rows = [
        row("example-hfl", "Q2 FY26", "aum", 10000.0), row("example-hfl", "Q1 FY26", "aum", 9600.0), row("example-hfl", "Q4 FY25", "aum", 9300.0),
        row("example-hfl", "Q2 FY26", "gnpa_pct", 1.82, status="needs_review", alt_value=1.2, alt_source="IP", footnote="QR 1.82% vs IP 1.20% differ by 34.07%, more than the 5% tolerance."),
        row("example-hfl", "Q2 FY26", "branches", 198, status="carried_forward", value_period="Q1 FY26", source="IP", footnote="Branches as of Q1 FY26 (previous quarter's IP)."),
        row("example-hfl", "Q2 FY26", "sell_down_volume", None, status="not_found", document=None, page_or_slide=None, footnote="AUM equals the loan book: no off-book loans."),
        row("example-hfl", "Q2 FY26", "aum", 10010.0, ts="2026-09-19T08:00:00Z", footnote="Re-extracted from the revised presentation."),
        row("sample-home-loans", "Q2 FY26", "aum", 4200.0),
    ]
    spec = build(rows, names={"example-hfl": "Example Housing Finance Ltd: [Standalone]/KPIs"})
    eq("sheets", [s["name"] for s in spec["sheets"]], ["Example Housing Finance Ltd   S", "sample-home-loans", "Footnotes", "Citations"])
    eq("sheet name <= 31", all(len(s["name"]) <= 31 for s in spec["sheets"]), True)
    s = spec["sheets"][0]
    eq("columns oldest quarter first", s["columns"], ["Category", "KPI", "Unit", "Q4 FY25", "Q1 FY26", "Q2 FY26", "Notes"])
    eq("27 KPI rows in catalog order", [r[1] for r in s["rows"]], [e["label"] for e in cat.CATALOG])
    eq("categories grouped in rulebook order", [c for i, c in enumerate(r[0] for r in s["rows"]) if i == 0 or c != s["rows"][i - 1][0]], cat.CATEGORIES)
    by = {r[1]: r for r in s["rows"]}
    eq("latest extraction wins", by["AUM"][3:6], [9300.0, 9600.0, 10010.0])
    eq("not_found cell text", by["Sell Down Volume"][5], "not found")
    eq("never-extracted cell is blank", by["CRAR %"][3:6], [None, None, None])
    eq("notes reference footnotes", (by["GNPA %"][6], by["Branches"][6]), ("Q2 FY26 [4] needs review", "Q2 FY26 [2] carried forward"))
    eq("cell status for colouring", sorted(c["status"] for c in s["cell_status"]), ["carried_forward", "needs_review", "not_found"])
    foot = spec["sheets"][2]["rows"]
    eq("footnote refs sequential", [f[0] for f in foot], [1, 2, 3, 4])
    eq("footnote carries value period", [f for f in foot if f[2] == "Branches"][0][5], "Q1 FY26")
    eq("citations for every latest cell", len(spec["sheets"][3]["rows"]), 7)
    eq("single-company filter", build(rows, customer="sample-home-loans")["filename"], "sample-home-loans-kpis.xlsx")
    used = {"footnotes"}
    eq("duplicate sheet names disambiguated", [sheet_name("Footnotes", used), sheet_name("Footnotes", used)], ["Footnotes (2)", "Footnotes (3)"])
    d = tempfile.mkdtemp(); p = os.path.join(d, "kpis.jsonl")
    with open(p, "w", encoding="utf-8") as f: f.write("\n".join(json.dumps(r, ensure_ascii=False) for r in rows) + "\n")
    out = run(p, xlsx=os.path.join(d, "out", "k.xlsx"))
    eq("run ok", out["ok"], True)
    eq("xlsx written iff openpyxl importable", out["xlsx"]["written"], os.path.exists(os.path.join(d, "out", "k.xlsx")))
    if not out["xlsx"]["written"]: eq("reason given when not written", "openpyxl" in (out["xlsx"]["reason"] or ""), True)
    with open(p, "a", encoding="utf-8") as f: f.write(json.dumps(row("example-hfl", "Q2 FY26", "nnpa_pct", 2.5), ensure_ascii=False) + "\n")
    out = run(p)
    eq("invalid input builds nothing", (out["ok"], out["errors"][0]["code"]), (False, "E-NNPA"))
    eq("unknown customer", run(p, customer="nobody")["ok"], False)
    return fails, 19


def main():
    ap = argparse.ArgumentParser(description="kpis.jsonl -> workbook spec JSON (one sheet per company, KPIs as rows by category, quarters as columns, Footnotes and Citations sheets) and, with --xlsx, the .xlsx.")
    ap.add_argument("path", nargs="?", help="kpis.jsonl with every row to show (history + the new batch)")
    ap.add_argument("--customer", help="build for this customer_id only")
    ap.add_argument("--name", action="append", default=[], metavar="CUSTOMER_ID=Display Name", help="sheet title for a company (repeatable); with --customer a bare name is accepted")
    ap.add_argument("--xlsx", help="write the workbook here (needs openpyxl, installed in the sandbox)")
    ap.add_argument("--self-test", action="store_true")
    a = ap.parse_args()
    if a.self_test:
        fails, n = _self_test()
        print(json.dumps({"script": "build_kpi_workbook", "cases": n, "failed": fails, "ok": not fails}, ensure_ascii=False, indent=1)); return 1 if fails else 0
    if not a.path: ap.print_help(sys.stderr); return 2
    names = {}
    for item in a.name:
        if "=" in item: k, v = item.split("=", 1); names[k.strip()] = v.strip()
        elif a.customer: names[a.customer] = item.strip()
        else: print(f"--name {item!r}: write CUSTOMER_ID=Display Name", file=sys.stderr); return 2
    try: out = run(a.path, a.customer, names, a.xlsx)
    except OSError as x: print(f"cannot read or write: {x}", file=sys.stderr); return 2
    print(json.dumps(out, ensure_ascii=False, indent=1))
    if not out["ok"]: print(out["problem"], file=sys.stderr); return 1
    if a.xlsx and not out["xlsx"]["written"]: print(out["xlsx"]["reason"], file=sys.stderr); return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
